from flask import Blueprint, request, jsonify, send_file
from ..middleware.auth_middleware import admin_required
from ..extensions import get_db
from datetime import date
import io
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

reports_bp = Blueprint('reports', __name__)

@reports_bp.route('/attendance/excel', methods=['GET'])
@admin_required
def export_attendance_excel():
    start = request.args.get('start', date.today().strftime('%Y-%m-%d'))
    end = request.args.get('end', date.today().strftime('%Y-%m-%d'))
    
    db = get_db()
    records = list(db.attendance.find({'date': {'$gte': start, '$lte': end}}).sort('date', 1))
    
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Attendance Report'
    
    # Header styling
    header_fill = PatternFill(start_color='4F46E5', end_color='4F46E5', fill_type='solid')
    header_font = Font(color='FFFFFF', bold=True, size=11)
    
    headers = ['Employee Name', 'Employee ID', 'Date', 'Check In', 'Check Out', 'Working Hours', 'Status', 'Late']
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal='center')
        ws.column_dimensions[cell.column_letter].width = 18
    
    for row, rec in enumerate(records, 2):
        check_in = rec.get('check_in')
        check_out = rec.get('check_out')
        ws.cell(row=row, column=1, value=rec.get('employee_name', ''))
        ws.cell(row=row, column=2, value=rec.get('employee_id', ''))
        ws.cell(row=row, column=3, value=rec.get('date', ''))
        ws.cell(row=row, column=4, value=check_in.strftime('%H:%M:%S') if check_in else '')
        ws.cell(row=row, column=5, value=check_out.strftime('%H:%M:%S') if check_out else 'Not checked out')
        ws.cell(row=row, column=6, value=round(rec.get('working_hours', 0), 2))
        ws.cell(row=row, column=7, value=rec.get('status', 'present'))
        ws.cell(row=row, column=8, value='Yes' if rec.get('is_late') else 'No')
    
    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    
    return send_file(
        buffer,
        as_attachment=True,
        download_name=f'attendance_{start}_to_{end}.xlsx',
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )

@reports_bp.route('/tasks/excel', methods=['GET'])
@admin_required
def export_tasks_excel():
    db = get_db()
    tasks = list(db.tasks.find().sort('created_at', -1))
    
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Tasks Report'
    
    header_fill = PatternFill(start_color='4F46E5', end_color='4F46E5', fill_type='solid')
    header_font = Font(color='FFFFFF', bold=True, size=11)
    
    headers = ['Task ID', 'Title', 'Priority', 'Status', 'Assigned By', 'Due Date', 'Created At', 'Completed At']
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal='center')
        ws.column_dimensions[cell.column_letter].width = 20
    
    for row, task in enumerate(tasks, 2):
        ws.cell(row=row, column=1, value=task.get('task_id', ''))
        ws.cell(row=row, column=2, value=task.get('title', ''))
        ws.cell(row=row, column=3, value=task.get('priority', ''))
        ws.cell(row=row, column=4, value=task.get('status', ''))
        ws.cell(row=row, column=5, value=task.get('assigned_by_name', ''))
        ws.cell(row=row, column=6, value=task.get('due_date', ''))
        ws.cell(row=row, column=7, value=task.get('created_at', '').strftime('%Y-%m-%d %H:%M') if task.get('created_at') else '')
        ws.cell(row=row, column=8, value=task.get('completed_at', '').strftime('%Y-%m-%d %H:%M') if task.get('completed_at') else '')
    
    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    
    return send_file(
        buffer,
        as_attachment=True,
        download_name=f'tasks_report.xlsx',
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
